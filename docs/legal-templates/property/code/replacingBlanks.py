from docx import Document
import re

input_file = r"E:\CODE\FYP\application\docs\legal-templates\property\docx\affidavitOfPossession.docx"
output_file = r"E:\CODE\FYP\application\docs\legal-templates\property\Placeholder-doc\affidavitOfPossession_with_placeholders.docx"


placeholders = [

# Deponent identity
"{{deponent_name}}",
"{{deponent_relation_name}}",
"{{deponent_age}}",
"{{deponent_cnic}}",
"{{deponent_profession}}",
"{{deponent_address_line1}}",
"{{deponent_address_line2}}",
"{{deponent_address_line3}}",

# Property details
"{{property_type}}",
"{{property_address}}",
"{{property_area_size}}",
"{{registry_document_number}}",
"{{other_property_details}}",

# Acquisition and possession
"{{property_acquisition_method}}",
"{{possession_since_date}}",

# Purpose / authority
"{{affidavit_purpose}}",
"{{authority_name}}",

# Signature section
"{{deponent_signature}}",
"{{deponent_name}}",
"{{deponent_cnic}}",
"{{affidavit_date}}",
"{{affidavit_place}}"

]

placeholder_index = 0


def replace_blanks_in_paragraph(paragraph):
    global placeholder_index

    pattern = r'_{3,}'

    if re.search(pattern, paragraph.text):
        new_text = paragraph.text

        while re.search(pattern, new_text) and placeholder_index < len(placeholders):
            new_text = re.sub(pattern, placeholders[placeholder_index], new_text, count=1)
            placeholder_index += 1

        paragraph.text = new_text


doc = Document(input_file)


for paragraph in doc.paragraphs:
    replace_blanks_in_paragraph(paragraph)


for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_blanks_in_paragraph(paragraph)


doc.save(output_file)

print("Done. File saved as:", output_file)